"""Integración del contexto administrativo con datos controlados de dos territorios."""
from datetime import timedelta
from contextlib import closing
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from agenda.models import Compromiso
from atenciones.models import AtencionCiudadana
from atenciones.permisos import queryset_atenciones
from auditoria.models import RegistroAuditoria
from auditoria.permisos import registros_autorizados
from catalogos.models import SolicitudCatalogoGeneral, TipoAtencion
from evidencias.models import Evidencia
from evidencias.permisos import puede_agregar_evidencia, puede_validar_evidencia
from sgr.aportes import aportes_visibles, puede_revisar
from sgr.forms import IndicadorFiltroForm, MetaForm
from sgr.models import AporteIndicador, Indicador, Meta
from sgr.test_ponderaciones import DatosPonderacionesMixin
from usuarios.models import Delegacion, Perfil, SolicitudTransferencia
from usuarios.sesion import CONTEXTO_GLOBAL, DELEGACION_SESION, get_delegacion_activa


class ContextoAdministradorTests(DatosPonderacionesMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.usuarios['administrador']
        self.usuario_b = get_user_model().objects.create_user('usuario_exclusivo_b')
        Perfil.objects.filter(usuario=self.usuario_b).update(
            rol=self.usuarios['funcionario'].perfil.rol, delegacion=self.b)
        self.usuario_b.perfil.delegaciones.set([self.b])
        self.act_a = self.crear_actividad('Actividad local A', self.a, self.usuarios['funcionario'])
        self.act_b = self.crear_actividad('Actividad exclusiva B', self.b, self.usuario_b)
        self.comp_a = self.crear_compromiso('Compromiso local A', self.a, self.usuarios['funcionario'])
        self.comp_b = self.crear_compromiso('Compromiso exclusivo B', self.b, self.usuario_b)
        self.ev_a = self.crear_evidencia(atencion=self.act_a)
        self.ev_b = self.crear_evidencia(atencion=self.act_b)
        self.ev_comp_a = self.crear_evidencia(compromiso=self.comp_a)
        self.ev_comp_b = self.crear_evidencia(compromiso=self.comp_b)
        self.tipo_a = TipoAtencion.objects.create(codigo='CTX-A', nombre='Tipo local A', es_general=False, delegacion=self.a)
        self.tipo_b = TipoAtencion.objects.create(codigo='CTX-B', nombre='Tipo exclusivo B', es_general=False, delegacion=self.b)
        self.entrar('administrador')

    def crear_actividad(self, titulo, delegacion, usuario):
        return AtencionCiudadana.objects.create(titulo=titulo, fecha=self.actividad.fecha,
            solicitante='Solicitante de prueba', tipo=self.tipo, subatencion=self.sub,
            descripcion='Prueba de contexto', delegacion=delegacion, responsable=usuario, creado_por=usuario)

    def crear_compromiso(self, titulo, delegacion, usuario):
        return Compromiso.objects.create(titulo=titulo, fecha_hora=timezone.now() + timedelta(days=2),
            delegacion=delegacion, responsable=usuario, creado_por=usuario, estado='pendiente')

    def crear_evidencia(self, **origen):
        return Evidencia.objects.create(**origen, archivo='contexto/prueba.jpg',
                                       subido_por=self.usuarios['funcionario'])

    def seleccionar(self, valor):
        return self.client.post(reverse('usuarios:delegacion_activa'), {'delegacion_id': valor})

    def request_contexto(self, usuario=None, valor=CONTEXTO_GLOBAL):
        request = RequestFactory().get('/')
        request.user = usuario or self.admin
        request.session = {DELEGACION_SESION: valor}
        return request

    def test_admin_selecciona_delegacion_y_muestra_contexto(self):
        self.assertEqual(self.seleccionar(self.a.pk).status_code, 302)
        self.assertEqual(self.client.session[DELEGACION_SESION], self.a.pk)
        respuesta = self.client.get(reverse('home'))
        self.assertContains(respuesta, str(self.a))
        self.assertEqual(respuesta.context['dashboard']['contexto_territorial'], self.a)

    def test_admin_vuelve_a_global_explicito(self):
        self.seleccionar(self.a.pk)
        self.assertEqual(self.seleccionar(CONTEXTO_GLOBAL).status_code, 302)
        self.assertEqual(self.client.session[DELEGACION_SESION], CONTEXTO_GLOBAL)
        respuesta = self.client.get(reverse('home'))
        self.assertIsNone(respuesta.context['sesion_delegacion'])
        self.assertContains(respuesta, 'Global: todas las delegaciones')
        self.assertEqual(respuesta.context['dashboard']['total_metas'], 2)

    def test_sesion_nueva_admin_es_global_sin_usar_primera_delegacion(self):
        request = self.request_contexto()
        request.session.clear()
        self.assertIsNone(get_delegacion_activa(request))
        self.assertEqual(set(queryset_atenciones(request).values_list('delegacion_id', flat=True)),
                         {self.a.pk, self.b.pk})

    def test_superusuario_tambien_usa_selector(self):
        user = get_user_model().objects.create_superuser('super-contexto', 'a@example.com', None)
        self.client.force_login(user)
        self.assertEqual(self.seleccionar(self.b.pk).status_code, 302)
        self.assertEqual(self.client.get(reverse('home')).context['dashboard']['contexto_territorial'], self.b)

    def test_persiste_al_navegar_por_los_modulos(self):
        self.seleccionar(self.a.pk)
        for ruta in ('home', 'atenciones:lista', 'agenda:lista', 'evidencias:lista',
                     'usuarios:lista', 'sgr:meta_lista', 'sgr:informe', 'auditoria:lista'):
            with self.subTest(ruta=ruta):
                respuesta = self.client.get(reverse(ruta))
                self.assertEqual(respuesta.status_code, 200)
                self.assertEqual(respuesta.context['sesion_delegacion'], self.a)
                self.assertEqual(self.client.session[DELEGACION_SESION], self.a.pk)

    def test_roles_territoriales_no_ven_ni_usan_selector(self):
        for rol in ('delegado', 'coordinador', 'funcionario', 'verificador'):
            with self.subTest(rol=rol):
                self.entrar(rol)
                self.assertNotContains(self.client.get(reverse('home')), 'id="selector-contexto-territorial"')
                for valor in (self.b.pk, self.a.pk, CONTEXTO_GLOBAL):
                    self.assertEqual(self.seleccionar(valor).status_code, 403)
                self.assertNotEqual(self.client.session.get(DELEGACION_SESION), self.b.pk)

    def test_usuario_sin_rol_no_puede_usar_selector(self):
        Perfil.objects.filter(usuario=self.usuarios['funcionario']).update(rol=None)
        self.entrar('funcionario')
        self.assertEqual(self.seleccionar(self.a.pk).status_code, 403)
        respuesta = self.client.get(reverse('home'))
        self.assertNotContains(respuesta, 'id="selector-contexto-territorial"')
        self.assertEqual(respuesta.context['total_actividades'], 0)

    def test_anonimo_redirigido_y_get_no_modifica_contexto(self):
        self.client.logout()
        self.assertEqual(self.seleccionar(self.a.pk).status_code, 302)
        self.entrar('administrador')
        respuesta = self.client.get(reverse('usuarios:delegacion_activa'), {'delegacion_id': self.a.pk})
        self.assertEqual(respuesta.status_code, 405)
        self.assertNotIn(DELEGACION_SESION, self.client.session)

    def test_rechaza_id_inexistente_desactivado_o_malformado_sin_ampliar_contexto(self):
        inactiva = Delegacion.objects.create(nombre='Delegación inactiva de prueba', activa=False)
        self.seleccionar(self.a.pk)
        for valor in ('', 'no-numero', '-1', '99999999', inactiva.pk, '1.5'):
            with self.subTest(valor=valor):
                self.assertEqual(self.seleccionar(valor).status_code, 400)
                self.assertEqual(self.client.session[DELEGACION_SESION], self.a.pk)

    def test_delegacion_desactivada_en_sesion_se_recupera_como_global(self):
        self.seleccionar(self.a.pk)
        self.a.activa = False
        self.a.save()
        respuesta = self.client.get(reverse('home'))
        self.assertIsNone(respuesta.context['sesion_delegacion'])
        self.assertContains(respuesta, 'Se restableció el contexto global')
        self.assertNotIn(DELEGACION_SESION, self.client.session)

    def test_delegacion_eliminada_en_sesion_se_recupera_sin_arbitraria(self):
        eliminable = Delegacion.objects.create(nombre='Delegación eliminable')
        self.seleccionar(eliminable.pk)
        eliminable.delete()
        self.assertIsNone(self.client.get(reverse('home')).context['sesion_delegacion'])

    def test_sesion_admin_invalida_no_produce_error(self):
        for valor in ('invalido', '99999999', [], True, 1.5):
            with self.subTest(valor=valor):
                request = self.request_contexto(valor=valor)
                self.assertIsNone(get_delegacion_activa(request))
                self.assertNotIn(DELEGACION_SESION, request.session)

    def test_contexto_se_resuelve_una_vez_por_solicitud(self):
        request = self.request_contexto(valor=self.a.pk)
        self.assertEqual(get_delegacion_activa(request), self.a)
        with self.assertNumQueries(0):
            for _ in range(5):
                self.assertEqual(get_delegacion_activa(request), self.a)
        request.session[DELEGACION_SESION] = self.b.pk
        self.assertEqual(get_delegacion_activa(request), self.b)

    def test_revocar_autorizacion_territorial_se_aplica_incluso_en_la_misma_solicitud(self):
        request = self.request_contexto(self.usuarios['funcionario'], self.a.pk)
        self.assertEqual(get_delegacion_activa(request), self.a)
        request.user.perfil.delegaciones.clear()
        self.assertIsNone(get_delegacion_activa(request))
        self.assertFalse(queryset_atenciones(request).exists())

    def test_detalles_no_muestran_relaciones_historicas_de_otro_territorio(self):
        self.comp_a.actividades.add(self.act_b)
        self.comp_b.actividades.add(self.act_a)
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('atenciones:detalle', args=[self.act_a.pk]))
        self.assertNotContains(respuesta, self.comp_b.titulo)
        respuesta = self.client.get(reverse('agenda:detalle', args=[self.comp_a.pk]))
        self.assertNotContains(respuesta, self.act_b.titulo)

    def test_parametros_y_sesion_ajena_no_amplian_a_roles_territoriales(self):
        for valor in (self.b.pk, CONTEXTO_GLOBAL):
            self.entrar('funcionario')
            sesion = self.client.session
            sesion[DELEGACION_SESION] = valor
            sesion.save()
            respuesta = self.client.get(reverse('atenciones:lista'),
                {'delegacion_id': self.b.pk, 'delegacion': CONTEXTO_GLOBAL})
            self.assertContains(respuesta, self.act_a.titulo)
            self.assertNotContains(respuesta, self.act_b.titulo)

    def test_csrf_obligatorio_y_cambio_valido_con_token(self):
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(self.admin)
        url = reverse('usuarios:delegacion_activa')
        self.assertEqual(cliente.post(url, {'delegacion_id': self.a.pk}).status_code, 403)
        cliente.get(reverse('home'))
        token = cliente.cookies['csrftoken'].value
        self.assertEqual(cliente.post(url, {'delegacion_id': self.a.pk, 'csrfmiddlewaretoken': token}).status_code, 302)

    def test_redireccion_externa_no_permitida_y_destino_local_conservado(self):
        url = reverse('usuarios:delegacion_activa')
        respuesta = self.client.post(url, {'delegacion_id': self.a.pk, 'next': 'https://example.org/'})
        self.assertRedirects(respuesta, reverse('home'), fetch_redirect_response=False)
        destino = reverse('sgr:informe') + '?periodo=' + str(self.periodo.pk)
        respuesta = self.client.post(url, {'delegacion_id': self.b.pk, 'next': destino})
        self.assertRedirects(respuesta, destino, fetch_redirect_response=False)

    def test_cambiar_contexto_no_modifica_registros_ni_auditoria(self):
        modelos = (AtencionCiudadana, Compromiso, Evidencia, Perfil, Meta, Indicador, RegistroAuditoria)
        antes = {m: list(m.objects.order_by('pk').values()) for m in modelos}
        for valor in (self.a.pk, self.b.pk, CONTEXTO_GLOBAL):
            self.seleccionar(valor)
        despues = {m: list(m.objects.order_by('pk').values()) for m in modelos}
        self.assertEqual(antes, despues)

    def test_listado_y_exportacion_actividades_respetan_contexto(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('atenciones:lista'), {'q': 'Actividad'})
        self.assertContains(respuesta, self.act_a.titulo)
        self.assertNotContains(respuesta, self.act_b.titulo)
        descarga = self.client.get(reverse('atenciones:lista'), {'export': 'xlsx'})
        with closing(load_workbook(BytesIO(descarga.content))) as libro:
            valores = [c.value for fila in libro.active for c in fila]
        self.assertIn(self.act_a.titulo, valores)
        self.assertNotIn(self.act_b.titulo, valores)

    def test_detalle_actividad_ajena_no_se_abre_por_id(self):
        self.seleccionar(self.a.pk)
        self.assertEqual(self.client.get(reverse('atenciones:detalle', args=[self.act_b.pk])).status_code, 404)
        self.seleccionar(CONTEXTO_GLOBAL)
        self.assertEqual(self.client.get(reverse('atenciones:detalle', args=[self.act_b.pk])).status_code, 200)

    def test_cambios_de_estado_y_disponibilidad_ajenos_se_rechazan(self):
        self.seleccionar(self.a.pk)
        for ruta, datos in (('atenciones:cambiar_estado', {'estado': 'en_proceso'}),
                            ('atenciones:cambiar_disponibilidad', {'activo': 'false'}),
                            ('atenciones:cambiar_va_al_tubo', {'va_al_tubo': 'true'})):
            self.assertEqual(self.client.post(reverse(ruta, args=[self.act_b.pk]), datos).status_code, 404)
        self.act_b.refresh_from_db()
        self.assertTrue(self.act_b.activo)
        self.assertEqual(self.act_b.estado, 'registrada')
        self.assertFalse(self.act_b.va_al_tubo)

    def test_cambio_de_estado_local_conserva_regla_y_territorio(self):
        self.seleccionar(self.a.pk)
        self.assertEqual(self.client.post(reverse('atenciones:cambiar_estado', args=[self.act_a.pk]),
                                         {'estado': 'en_proceso'}).status_code, 302)
        self.act_a.refresh_from_db()
        self.assertEqual(self.act_a.estado, 'en_proceso')
        self.assertEqual(self.act_a.delegacion_id, self.a.pk)

    def test_crear_actividad_global_no_elige_delegacion_arbitraria(self):
        self.seleccionar(CONTEXTO_GLOBAL)
        antes = AtencionCiudadana.objects.count()
        respuesta = self.client.post(reverse('atenciones:crear'), self.datos_actividad())
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(AtencionCiudadana.objects.count(), antes)

    def test_crear_actividad_local_ignora_delegacion_inyectada(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.post(reverse('atenciones:crear'), self.datos_actividad(
            titulo='Alta en contexto A', delegacion=self.b.pk))
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(AtencionCiudadana.objects.get(titulo='Alta en contexto A').delegacion_id, self.a.pk)

    def test_agenda_lista_resumen_y_opciones_de_filtro_son_territoriales(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('agenda:lista'))
        self.assertContains(respuesta, self.comp_a.titulo)
        self.assertNotContains(respuesta, self.comp_b.titulo)
        self.assertNotIn(self.usuario_b, respuesta.context['filtro'].fields['responsable'].queryset)
        self.assertTrue(all(c.delegacion_id == self.a.pk for c in respuesta.context['proximos']))

    def test_compromiso_ajeno_no_admite_detalle_edicion_ni_asociacion(self):
        self.seleccionar(self.a.pk)
        for ruta in ('agenda:detalle', 'agenda:editar'):
            self.assertEqual(self.client.get(reverse(ruta, args=[self.comp_b.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('agenda:editar', args=[self.comp_b.pk]), {}).status_code, 404)
        self.assertEqual(self.client.get(reverse('atenciones:agregar_a_agenda', args=[self.act_b.pk])).status_code, 404)

    def test_formulario_compromiso_solo_ofrece_actividades_de_su_origen(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('agenda:editar', args=[self.comp_a.pk]))
        opciones = respuesta.context['form'].fields['actividades'].queryset
        self.assertIn(self.act_a, opciones)
        self.assertNotIn(self.act_b, opciones)

    def test_evidencias_de_ambos_origenes_respetan_contexto(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('evidencias:lista'))
        self.assertEqual(set(respuesta.context['evidencias'].values_list('pk', flat=True)),
                         {self.ev_a.pk, self.ev_comp_a.pk})
        for ev in (self.ev_b, self.ev_comp_b):
            self.assertEqual(self.client.get(reverse('evidencias:detalle', args=[ev.pk])).status_code, 404)
            self.assertEqual(self.client.get(reverse('evidencia_archivo', args=[ev.pk])).status_code, 404)

    def test_subida_ajena_se_bloquea_en_vista_y_permiso(self):
        self.seleccionar(self.a.pk)
        request = self.request_contexto(valor=self.a.pk)
        self.assertFalse(puede_agregar_evidencia(request, self.act_b))
        self.assertFalse(puede_agregar_evidencia(request, self.comp_b))
        self.assertEqual(self.client.get(reverse('evidencias:atencion_crear', args=[self.act_b.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('evidencias:agenda_crear', args=[self.comp_b.pk])).status_code, 404)

    def test_formulario_evidencia_no_ofrece_origen_ajeno(self):
        self.seleccionar(self.a.pk)
        form = self.client.get(reverse('evidencias:crear')).context['form']
        self.assertNotIn(self.act_b, form.fields['atencion'].queryset)
        self.assertNotIn(self.comp_b, form.fields['compromiso'].queryset)

    def test_validacion_ajena_no_es_posible_por_post(self):
        self.seleccionar(self.a.pk)
        for evidencia in (self.ev_b, self.ev_comp_b):
            respuesta = self.client.post(reverse('evidencias:validar', args=[evidencia.pk]), {'accion': 'aprobar'})
            self.assertEqual(respuesta.status_code, 404)
            evidencia.refresh_from_db()
            self.assertEqual(evidencia.estado, 'pendiente')

    def test_admin_no_valida_evidencia_de_actividad_propia_en_ningun_contexto(self):
        self.act_a.creado_por = self.admin
        self.act_a.save(update_fields=['creado_por'])
        for valor in (self.a.pk, CONTEXTO_GLOBAL):
            self.seleccionar(valor)
            self.assertFalse(puede_validar_evidencia(self.request_contexto(valor=valor), self.ev_a))
            self.assertEqual(self.client.post(reverse('evidencias:validar', args=[self.ev_a.pk]),
                                             {'accion': 'aprobar'}).status_code, 404)

    def test_admin_no_valida_evidencia_de_compromiso_propio(self):
        self.comp_a.creado_por = self.admin
        self.comp_a.save(update_fields=['creado_por'])
        self.assertFalse(puede_validar_evidencia(self.request_contexto(), self.ev_comp_a))

    def test_admin_conserva_validacion_de_evidencia_ajena_en_contexto(self):
        self.seleccionar(self.a.pk)
        self.assertEqual(self.client.post(reverse('evidencias:validar', args=[self.ev_a.pk]),
                                         {'accion': 'aprobar'}).status_code, 302)
        self.ev_a.refresh_from_db()
        self.assertEqual(self.ev_a.estado, 'aprobada')

    def test_usuarios_listado_resumen_y_exportacion_respetan_autorizaciones(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('usuarios:lista'))
        self.assertNotContains(respuesta, self.usuario_b.username)
        self.assertNotIn(self.usuario_b.pk, [p.usuario_id for p in respuesta.context['perfiles']])
        esperado = Perfil.objects.filter(delegaciones=self.a).count()
        self.assertEqual(sum(f['total'] for f in respuesta.context['resumen_roles']), esperado)
        descarga = self.client.get(reverse('usuarios:lista'), {'export': 'xlsx'})
        with closing(load_workbook(BytesIO(descarga.content))) as libro:
            self.assertNotIn(self.usuario_b.username, [c.value for fila in libro.active for c in fila])

    def test_usuarios_ajenos_no_admiten_edicion_o_desactivacion(self):
        self.seleccionar(self.a.pk)
        perfil = self.usuario_b.perfil
        for ruta in ('usuarios:detalle', 'usuarios:editar'):
            self.assertEqual(self.client.get(reverse(ruta, args=[perfil.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('usuarios:desactivar', args=[perfil.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('usuarios:editar', args=[perfil.pk]), {}).status_code, 404)

    def test_alta_usuario_no_admite_autorizaciones_fuera_de_contexto(self):
        self.seleccionar(self.a.pk)
        form = self.client.get(reverse('usuarios:crear')).context['form']
        self.assertIn(self.a, form.fields['delegaciones'].queryset)
        self.assertNotIn(self.b, form.fields['delegaciones'].queryset)
        respuesta = self.client.post(reverse('usuarios:crear'), {
            'username': 'alta_ajena_inyectada', 'first_name': 'Prueba', 'last_name': 'Local',
            'email': 'alta@example.com', 'rol': self.usuarios['funcionario'].perfil.rol_id,
            'delegaciones': [self.b.pk], 'password1': 'ClavePruebaSegura2026!',
            'password2': 'ClavePruebaSegura2026!', 'activo': 'on',
        })
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('delegaciones', respuesta.context['form'].errors)
        self.assertFalse(get_user_model().objects.filter(username='alta_ajena_inyectada').exists())

    def test_usuario_autorizado_en_dos_delegaciones_se_incluye_sin_duplicarlo(self):
        self.usuario_b.perfil.delegaciones.add(self.a)
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('usuarios:lista'))
        self.assertEqual(sum(p.usuario_id == self.usuario_b.pk for p in respuesta.context['perfiles']), 1)

    def test_metas_indicadores_y_filtros_respetan_contexto(self):
        self.seleccionar(self.a.pk)
        for ruta, local, ajeno in (('sgr:meta_lista', self.meta_a.nombre, self.meta_b.nombre),
                                   ('sgr:indicador_lista', self.indicador_a.nombre, self.indicador_b.nombre)):
            respuesta = self.client.get(reverse(ruta))
            self.assertContains(respuesta, local)
            self.assertNotContains(respuesta, ajeno)
        request = self.request_contexto(valor=self.a.pk)
        opciones = IndicadorFiltroForm(request=request).fields['meta'].queryset
        self.assertNotIn(self.meta_b, opciones)

    def test_parametro_filtro_sgr_ajeno_no_amplia_contexto_admin(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('sgr:meta_lista'), {'delegacion': self.b.pk})
        self.assertNotContains(respuesta, self.meta_b.nombre)
        self.assertFalse(respuesta.context['form'].is_valid())

    def test_detalles_y_escrituras_sgr_ajenos_se_bloquean(self):
        self.seleccionar(self.a.pk)
        for ruta, pk in (('sgr:meta_detalle', self.meta_b.pk), ('sgr:meta_editar', self.meta_b.pk),
                          ('sgr:indicador_detalle', self.indicador_b.pk), ('sgr:indicador_editar', self.indicador_b.pk)):
            self.assertEqual(self.client.get(reverse(ruta, args=[pk])).status_code, 404)
        for ruta, pk in (('sgr:meta_estado', self.meta_b.pk), ('sgr:indicador_estado', self.indicador_b.pk)):
            self.assertEqual(self.client.post(reverse(ruta, args=[pk]), {'activo': 'false'}).status_code, 404)

    def test_creacion_meta_no_admite_delegacion_fuera_de_contexto(self):
        request = self.request_contexto(valor=self.a.pk)
        form = MetaForm(self.datos(nombre='Meta inyectada', delegacion=self.b.pk), request=request)
        self.assertFalse(form.is_valid())
        self.assertIn('delegacion', form.errors)

    def test_dashboard_y_contadores_son_coherentes_con_contexto(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('home'))
        self.assertEqual(respuesta.context['total_actividades'], AtencionCiudadana.objects.filter(delegacion=self.a).count())
        self.assertEqual(respuesta.context['compromisos_proximos'], 1)
        self.assertEqual(respuesta.context['evidencias_pendientes'], 2)
        self.assertEqual(respuesta.context['total_usuarios'], Perfil.objects.filter(activo=True, delegaciones=self.a).count())
        self.assertEqual([f.meta.pk for f in respuesta.context['dashboard']['filas']], [self.meta_a.pk])

    def test_informe_dashboard_y_xlsx_usan_mismo_contexto_y_resultados(self):
        self.evaluado('100', '120')
        self.seleccionar(self.a.pk)
        dashboard = self.client.get(reverse('home')).context['dashboard']
        respuesta = self.client.get(reverse('sgr:informe'), {'delegacion': self.b.pk})
        informe = respuesta.context['informe']
        self.assertEqual(informe['contexto_territorial'], self.a)
        self.assertEqual(informe['filas'][0].resultado, dashboard['filas'][0].resultado)
        self.assertNotContains(respuesta, self.meta_b.nombre)
        descarga = self.client.get(reverse('sgr:informe'), {'export': 'xlsx', 'delegacion': self.b.pk})
        with closing(load_workbook(BytesIO(descarga.content))) as libro:
            self.assertEqual(libro['Metas']['B3'].value, str(self.a))
            self.assertEqual([f[0] for f in libro['Metas'].iter_rows(min_row=8, values_only=True)], [self.meta_a.nombre])

    def test_solo_cambia_alcance_sin_alterar_formulas_al_volver_a_global(self):
        self.evaluado('100', '231')
        self.seleccionar(self.a.pk)
        local = self.client.get(reverse('sgr:informe')).context['informe']['filas'][0].resultado
        self.seleccionar(CONTEXTO_GLOBAL)
        filas = self.client.get(reverse('sgr:informe')).context['informe']['filas']
        global_a = next(f.resultado for f in filas if f.meta.pk == self.meta_a.pk)
        self.assertEqual(local, global_a)

    def test_aportes_y_permisos_de_revision_respetan_contexto(self):
        aporte_b = self.vincular('cerrada', AporteIndicador.Estado.PENDIENTE, indicador=self.indicador_b)
        request = self.request_contexto(valor=self.a.pk)
        self.assertNotIn(aporte_b, aportes_visibles(request))
        self.assertFalse(puede_revisar(request, aporte_b))
        self.assertTrue(puede_revisar(self.request_contexto(), aporte_b))

    def test_auditoria_se_filtra_por_objeto_y_sus_opciones_no_filtran_otro_territorio(self):
        for objeto, etiqueta, actor in ((self.act_a, 'Evento local A', self.admin),
                                       (self.act_b, 'Evento exclusivo B', self.usuario_b)):
            RegistroAuditoria.objects.create(usuario=actor, accion='crear', modelo=objeto._meta.label,
                                             objeto_id=str(objeto.pk), descripcion=etiqueta)
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('auditoria:lista'))
        self.assertContains(respuesta, 'Evento local A')
        self.assertNotContains(respuesta, 'Evento exclusivo B')
        self.assertNotIn(self.usuario_b, respuesta.context['form'].fields['usuario'].queryset)
        self.seleccionar(CONTEXTO_GLOBAL)
        self.assertContains(self.client.get(reverse('auditoria:lista')), 'Evento exclusivo B')

    def test_auditoria_de_objetos_generales_solo_es_global(self):
        registro = RegistroAuditoria.objects.create(usuario=self.admin, accion='editar', modelo='sgr.Periodo',
            objeto_id=str(self.periodo.pk), descripcion='Configuración de período global')
        qs = RegistroAuditoria.objects.filter(pk=registro.pk)
        self.assertFalse(registros_autorizados(self.request_contexto(valor=self.a.pk), qs).exists())
        self.assertTrue(registros_autorizados(self.request_contexto(), qs).exists())

    def test_auditoria_de_usuarios_admin_se_asocia_por_autorizaciones(self):
        registro = RegistroAuditoria.objects.create(usuario=self.admin, accion='editar', modelo='auth.User',
            objeto_id=str(self.usuario_b.pk), descripcion='Edición usuario B')
        qs = RegistroAuditoria.objects.filter(pk=registro.pk)
        self.assertFalse(registros_autorizados(self.request_contexto(valor=self.a.pk), qs).exists())
        self.assertTrue(registros_autorizados(self.request_contexto(valor=self.b.pk), qs).exists())

    def test_catalogos_globales_se_consultan_y_locales_ajenos_no(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.get(reverse('catalogos:tipo_lista'))
        self.assertContains(respuesta, self.tipo.nombre)
        self.assertContains(respuesta, self.tipo_a.nombre)
        self.assertNotContains(respuesta, self.tipo_b.nombre)

    def test_catalogos_locales_ajenos_no_se_editan_y_los_generales_se_administran_en_global(self):
        self.seleccionar(self.a.pk)
        for tipo in (self.tipo_b, self.tipo):
            self.assertEqual(self.client.get(reverse('catalogos:tipo_editar', args=[tipo.pk])).status_code, 404)
        self.seleccionar(CONTEXTO_GLOBAL)
        self.assertEqual(self.client.get(reverse('catalogos:tipo_editar', args=[self.tipo.pk])).status_code, 200)

    def test_transferencia_directa_restringe_origen_pero_no_reasigna_registros(self):
        self.seleccionar(self.a.pk)
        respuesta = self.client.post(reverse('usuarios:transferencia_directa', args=[self.usuario_b.pk]),
            {'delegacion_destino': self.a.pk, 'confirmar_compromisos': 'on'})
        self.assertEqual(respuesta.status_code, 403)
        self.usuario_b.perfil.refresh_from_db()
        self.assertEqual(self.usuario_b.perfil.delegacion_id, self.b.pk)
        self.comp_b.refresh_from_db()
        self.assertEqual(self.comp_b.delegacion_id, self.b.pk)

    def crear_solicitudes_propias(self):
        usuario = self.usuarios['coordinador']
        SolicitudTransferencia.objects.create(usuario=self.usuarios['funcionario'],
            delegacion_origen=self.a, delegacion_destino=self.b, solicitante=usuario)
        SolicitudCatalogoGeneral.objects.create(tipo_solicitud='categoria', codigo='CTX-SOL',
            nombre='Solicitud propia A', delegacion_solicitante=self.a, solicitante=usuario)
        return usuario

    def test_solicitudes_propias_no_se_muestran_sin_autorizacion_territorial(self):
        usuario = self.crear_solicitudes_propias()
        usuario.perfil.delegaciones.clear()
        self.entrar('coordinador')
        for ruta in ('usuarios:transferencias', 'catalogos:solicitud_lista'):
            respuesta = self.client.get(reverse(ruta))
            self.assertEqual(respuesta.status_code, 200)
            self.assertFalse(respuesta.context['solicitudes'].exists())

    def test_solicitudes_propias_fuera_de_delegacion_activa_no_se_muestran(self):
        usuario = self.crear_solicitudes_propias()
        usuario.perfil.delegaciones.add(self.b)
        self.entrar('coordinador')
        sesion = self.client.session
        sesion[DELEGACION_SESION] = self.b.pk
        sesion.save()
        for ruta in ('usuarios:transferencias', 'catalogos:solicitud_lista'):
            self.assertFalse(self.client.get(reverse(ruta)).context['solicitudes'].exists())

    def test_delegado_sin_contexto_no_recibe_perfiles_sin_delegacion(self):
        get_user_model().objects.create_user('perfil_sin_delegacion')
        self.usuarios['delegado'].perfil.delegaciones.clear()
        self.entrar('delegado')
        respuesta = self.client.get(reverse('usuarios:transferencias'))
        self.assertFalse(respuesta.context['perfiles_transferibles'].exists())

    def test_aportes_propios_no_se_validan_en_global_ni_territorial(self):
        indicador = self.nuevo('100')
        aporte = self.vincular('cerrada', AporteIndicador.Estado.PENDIENTE, indicador=indicador)
        aporte.actividad.creado_por = self.admin
        aporte.actividad.save(update_fields=['creado_por'])
        for valor in (self.a.pk, CONTEXTO_GLOBAL):
            self.assertFalse(puede_revisar(self.request_contexto(valor=valor), aporte))

    def test_roles_territoriales_no_reciben_permisos_de_mantenimiento_sgr(self):
        for rol in ('funcionario', 'verificador'):
            self.entrar(rol)
            self.assertEqual(self.client.post(reverse('sgr:meta_crear'), self.datos()).status_code, 403)
            self.assertEqual(self.seleccionar(CONTEXTO_GLOBAL).status_code, 403)

    def test_usuario_territorial_sin_autorizacion_no_ve_datos(self):
        self.usuarios['funcionario'].perfil.delegaciones.clear()
        self.entrar('funcionario')
        for ruta in ('atenciones:lista', 'agenda:lista', 'evidencias:lista'):
            respuesta = self.client.get(reverse(ruta))
            self.assertNotContains(respuesta, self.act_a.titulo)
            self.assertNotContains(respuesta, self.comp_a.titulo)
        self.assertEqual(self.client.get(reverse('sgr:informe')).status_code, 403)

    def test_lecturas_en_contexto_no_escriben_registros(self):
        self.seleccionar(self.a.pk)
        with CaptureQueriesContext(connection) as consultas:
            for ruta in ('home', 'sgr:informe', 'atenciones:lista', 'agenda:lista'):
                self.client.get(reverse(ruta))
        escrituras = [q['sql'] for q in consultas if q['sql'].lstrip().split()[0].upper()
                      in ('INSERT', 'UPDATE', 'DELETE', 'REPLACE')]
        self.assertEqual(escrituras, [])
