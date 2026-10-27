from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.db import connection
from django.test import RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from sgr.calculos import calcular_resultado_indicador
from sgr.exportacion_informes import exportar_informe_excel
from sgr.informes import construir_informe
from sgr.models import AporteIndicador, Indicador, Meta, Periodo
from sgr.test_ponderaciones import DatosPonderacionesMixin
from usuarios.models import Perfil


class InformeSGRTests(DatosPonderacionesMixin, TestCase):
    def setUp(self):
        super().setUp()
        fecha = patch('sgr.informes.timezone.localdate', return_value=date(2026, 10, 26))
        fecha.start()
        self.addCleanup(fecha.stop)
        self.entrar('administrador')

    def consultar(self, **parametros):
        return self.client.get(reverse('sgr:informe'), parametros)

    def informe(self, **parametros):
        return self.consultar(**parametros).context['informe']

    def fila_a(self):
        return next(f for f in self.informe()['filas'] if f.meta.pk == self.meta_a.pk)

    def libro(self, **parametros):
        respuesta = self.consultar(export='xlsx', **parametros)
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        libro = load_workbook(BytesIO(respuesta.content))
        self.addCleanup(libro.close)
        return libro

    def test_anonimo_redirigido_a_login_tambien_al_exportar(self):
        self.client.logout()
        for parametros in ({}, {'export': 'xlsx'}):
            respuesta = self.consultar(**parametros)
            self.assertEqual(respuesta.status_code, 302)
            self.assertTrue(respuesta.url.startswith(reverse('login') + '?next='))

    def test_roles_validos_consultan_y_ven_navegacion(self):
        for rol in ('administrador', 'delegado', 'coordinador', 'funcionario', 'verificador'):
            with self.subTest(rol=rol):
                self.entrar(rol)
                self.assertEqual(self.consultar().status_code, 200)
                self.assertContains(self.client.get(reverse('home')), f'href="{reverse("sgr:informe")}"')

    def test_usuario_sin_rol_no_obtiene_html_ni_exportacion(self):
        usuario = get_user_model().objects.create_user('sin-rol')
        Perfil.objects.filter(usuario=usuario).update(rol=None)
        self.client.force_login(usuario)
        for parametros in ({}, {'export': 'xlsx'}, {'periodo': 'invalido'}):
            self.assertEqual(self.consultar(**parametros).status_code, 403)
        self.assertNotContains(self.client.get(reverse('home')), f'href="{reverse("sgr:informe")}"')

    def test_perfil_y_rol_inactivos_no_consultan(self):
        usuario = self.usuarios['funcionario']
        for campo in ('perfil', 'rol'):
            with self.subTest(campo=campo):
                Perfil.objects.filter(pk=usuario.perfil.pk).update(activo=campo != 'perfil')
                rol = usuario.perfil.rol
                rol.activo = campo != 'rol'
                rol.save()
                self.entrar('funcionario')
                respuesta = self.consultar()
                if campo == 'perfil':
                    self.assertRedirects(respuesta, reverse('login') + '?next=' + reverse('sgr:informe'),
                                         fetch_redirect_response=False)
                else:
                    self.assertEqual(respuesta.status_code, 403)

    def test_sin_delegacion_autorizada_no_consulta_ni_exporta(self):
        self.usuarios['delegado'].perfil.delegaciones.clear()
        self.entrar('delegado')
        for parametros in ({}, {'export': 'xlsx'}):
            respuesta = self.consultar(**parametros)
            self.assertEqual(respuesta.status_code, 403)
            self.assertNotContains(respuesta, self.meta_b.nombre, status_code=403)

    def test_delegacion_inactiva_no_consulta(self):
        self.a.activa = False
        self.a.save()
        self.entrar('verificador')
        self.assertEqual(self.consultar().status_code, 403)

    def test_servicio_autoriza_sin_depender_de_la_vista(self):
        request = RequestFactory().get('/')
        request.user = get_user_model().objects.create_user('sin-rol-servicio')
        request.session = {}
        with self.assertRaises(PermissionDenied):
            construir_informe(request, self.periodo)

    def test_administrador_perfil_ve_ambas_delegaciones(self):
        informe = self.informe()
        self.assertEqual({f.meta.pk for f in informe['filas']}, {self.meta_a.pk, self.meta_b.pk})
        self.assertEqual(informe['contexto_territorial'], 'Global: todas las delegaciones')

    def test_superusuario_ve_global_sin_delegacion_de_perfil(self):
        usuario = get_user_model().objects.create_superuser('super-informe', 'a@example.com', 'clave')
        self.client.force_login(usuario)
        self.assertEqual(self.informe()['total_metas'], 2)

    def test_cuatro_roles_territoriales_solo_ven_su_informacion(self):
        for rol in ('delegado', 'coordinador', 'funcionario', 'verificador'):
            with self.subTest(rol=rol):
                self.entrar(rol)
                respuesta = self.consultar()
                self.assertContains(respuesta, self.meta_a.nombre)
                self.assertNotContains(respuesta, self.meta_b.nombre)
                self.assertNotContains(respuesta, self.indicador_b.nombre)
                self.assertEqual(respuesta.context['informe']['total_metas'], 1)

    def test_parametros_territoriales_y_de_objetos_no_amplian_alcance(self):
        self.entrar('funcionario')
        respuesta = self.consultar(delegacion=self.b.pk, meta=self.meta_b.pk,
                                   indicador=self.indicador_b.pk, periodo=self.periodo.pk)
        self.assertContains(respuesta, self.meta_a.nombre)
        self.assertNotContains(respuesta, self.meta_b.nombre)
        self.assertNotContains(respuesta, self.indicador_b.nombre)

    def test_sesion_manipulada_no_autoriza_otra_delegacion(self):
        self.entrar('delegado')
        sesion = self.client.session
        sesion['delegacion_activa_id'] = self.b.pk
        sesion.save()
        self.assertEqual([f.meta.pk for f in self.informe()['filas']], [self.meta_a.pk])

    def test_enlaces_de_detalle_respetan_permisos_y_no_hay_acciones_de_edicion(self):
        indicador = self.nuevo('100')
        for rol in ('administrador', 'delegado', 'coordinador', 'funcionario', 'verificador'):
            with self.subTest(rol=rol):
                self.entrar(rol)
                respuesta = self.consultar()
                for enlace in (self.meta_a.get_absolute_url(), indicador.get_absolute_url()):
                    if rol in ('administrador', 'delegado', 'coordinador'):
                        self.assertContains(respuesta, f'href="{enlace}"')
                    else:
                        self.assertNotContains(respuesta, f'href="{enlace}"')
                        self.assertEqual(self.client.get(enlace).status_code, 403)
                self.assertNotContains(respuesta, reverse('sgr:indicador_editar', args=[indicador.pk]))
                self.assertNotContains(respuesta, reverse('sgr:indicador_estado', args=[indicador.pk]))

    def test_periodo_vigente_fechas_y_contexto_visibles(self):
        respuesta = self.consultar()
        self.assertEqual(respuesta.context['informe']['periodo'], self.periodo)
        for texto in (self.periodo.nombre, '01/01/2026', '31/12/2026', '26/10/2026',
                      'Global: todas las delegaciones', 'Resultado ponderado (pp)'):
            self.assertContains(respuesta, texto)

    def test_sin_vigente_no_se_sustituye_por_futuro_y_no_descarga_libro(self):
        Periodo.objects.filter(pk=self.periodo.pk).update(activo=False)
        respuesta = self.consultar(export='xlsx')
        self.assertContains(respuesta, 'No hay un período activo vigente')
        self.assertIsNone(respuesta.context['informe']['periodo'])
        self.assertNotIn('Content-Disposition', respuesta)

    def test_periodos_conflictivos_no_se_eligen_arbitrariamente(self):
        Periodo.objects.filter(pk=self.otro_periodo.pk).update(
            fecha_inicio=date(2026, 10, 1), fecha_fin=date(2026, 11, 1))
        respuesta = self.consultar()
        self.assertContains(respuesta, 'fechas inválidas o superpuestas')
        self.assertEqual(respuesta.context['informe']['filas'], [])

    def test_seleccion_explicita_de_periodo_inactivo_historico(self):
        self.evaluado('100', '120')
        Periodo.objects.filter(pk=self.periodo.pk).update(activo=False)
        with patch('sgr.informes.timezone.localdate', return_value=date(2028, 1, 1)):
            respuesta = self.consultar(periodo=self.periodo.pk)
        informe = respuesta.context['informe']
        self.assertEqual(informe['periodo'].pk, self.periodo.pk)
        self.assertEqual(informe['esperado'], Decimal('100'))
        fila = next(f for f in informe['filas'] if f.meta.pk == self.meta_a.pk)
        self.assertEqual(fila.resultado.cumplimiento_normalizado, Decimal('120'))
        self.assertContains(respuesta, 'no reconstruye estados')
        self.assertContains(respuesta, 'Período inactivo')

    def test_periodo_futuro_se_consulta_sin_asignar_color(self):
        self.crear_meta(periodo=self.otro_periodo, nombre='Meta de 2027')
        informe = self.informe(periodo=self.otro_periodo.pk)
        self.assertEqual(informe['esperado'], Decimal('0'))
        self.assertEqual(informe['filas'][0].semaforo.estado, 'sin_datos')

    def test_periodo_invalido_no_devuelve_datos_de_otro_periodo(self):
        for valor in ('texto', '-1', '99999999'):
            with self.subTest(valor=valor):
                respuesta = self.consultar(periodo=valor, export='xlsx')
                self.assertContains(respuesta, 'Seleccione un período válido', status_code=400)
                self.assertIsNone(respuesta.context['informe'])
                self.assertNotContains(respuesta, self.meta_a.nombre, status_code=400)
                self.assertNotIn('Content-Disposition', respuesta)

    def test_sin_metas_activas_es_informativo(self):
        Meta.objects.filter(periodo=self.periodo).update(activo=False)
        self.assertContains(self.consultar(), 'No hay metas activas para el período')
        self.assertEqual(self.informe()['total_metas'], 0)

    def test_sin_indicadores_conserva_meta_sin_resultado(self):
        meta = self.crear_meta(nombre='Meta sin indicadores')
        self.entrar('delegado')
        fila = next(f for f in self.informe()['filas'] if f.meta.pk == meta.pk)
        self.assertEqual(fila.resultado.indicadores, ())
        self.assertIsNone(fila.resultado.cumplimiento_normalizado)
        self.assertEqual(fila.semaforo.estado, 'sin_datos')
        self.assertContains(self.consultar(), 'No hay Indicadores para mostrar.')

    def test_metas_inactivas_y_de_otro_periodo_se_excluyen(self):
        Meta.objects.filter(pk=self.meta_a.pk).update(activo=False)
        self.crear_meta(periodo=self.otro_periodo, nombre='Meta futura')
        self.assertEqual([f.meta.pk for f in self.informe()['filas']], [self.meta_b.pk])

    def test_cero_real_y_none_son_distintos(self):
        cero = self.nuevo('40', unidad_medida='cantidad', unidad_personalizada='')
        sin_datos = self.evaluado('60', None)
        fila = self.fila_a()
        resultados = {p.indicador.pk: p for p in fila.resultado.indicadores}
        self.assertEqual(resultados[cero.pk].metrica.resultado, Decimal('0'))
        self.assertEqual(resultados[cero.pk].resultado_ponderado, Decimal('0'))
        self.assertIsNone(resultados[sin_datos.pk].metrica.resultado)
        self.assertIsNone(resultados[sin_datos.pk].resultado_ponderado)
        self.assertEqual(fila.resultado.cobertura_calculable, Decimal('40'))
        self.assertEqual(fila.resultado.cumplimiento_normalizado, Decimal('0'))
        self.assertEqual(fila.semaforo.estado, 'rojo')

    def test_sin_limite_no_se_convierte_en_cero(self):
        self.nuevo('100', sentido='menor', unidad_medida='cantidad', unidad_personalizada='')
        fila = self.fila_a()
        self.assertEqual(fila.resultado.sin_limite, 1)
        self.assertIsNone(fila.resultado.cumplimiento_normalizado)
        self.assertEqual(fila.semaforo.estado, 'sin_datos')
        self.assertContains(self.consultar(), 'Superior al objetivo (resultado 0), sin porcentaje numérico')

    def test_sobrecumplimiento_y_ponderacion_completa(self):
        self.evaluado('100', '231')
        fila = self.fila_a()
        self.assertEqual(fila.resultado.cumplimiento_normalizado, Decimal('231'))
        self.assertEqual(fila.resultado.resultado_ponderado_bruto, Decimal('231'))
        self.assertTrue(fila.resultado.configuracion_completa)
        self.assertEqual(fila.semaforo.estado, 'verde')
        self.assertContains(self.consultar(), '231,00%')

    def test_cobertura_parcial_y_pesos_sin_configurar_no_se_ocultan(self):
        self.evaluado('40', '120')
        self.evaluado('30', None)
        self.nuevo(None)
        fila = self.fila_a()
        self.assertEqual(fila.resultado.cobertura_calculable, Decimal('40'))
        self.assertEqual(fila.resultado.resultado_ponderado_bruto, Decimal('48'))
        self.assertEqual(fila.resultado.cumplimiento_normalizado, Decimal('120'))
        self.assertEqual(fila.resultado.sin_ponderacion, 1)
        self.assertFalse(fila.resultado.configuracion_completa)
        for texto in ('Cobertura parcial', 'Configuración incompleta', 'Ponderación no configurada'):
            self.assertContains(self.consultar(), texto)

    def test_inactivos_se_identifican_y_no_participan_en_cobertura(self):
        self.evaluado('40', '120')
        historico = self.nuevo('60', activo=False)
        fila = self.fila_a()
        ponderado = next(p for p in fila.resultado.indicadores if p.indicador.pk == historico.pk)
        self.assertEqual(ponderado.estado, 'inactivo')
        self.assertIsNone(ponderado.resultado_ponderado)
        self.assertEqual(fila.resultado.cobertura_calculable, Decimal('40'))
        self.assertContains(self.consultar(), 'Excluido: inactivo')

    def test_reporte_detalle_dashboard_y_metrica_individual_coinciden(self):
        indicador = self.evaluado('100', '120')
        self.entrar('delegado')
        fila = self.fila_a()
        detalle = self.client.get(self.meta_a.get_absolute_url()).context['resultado_meta']
        dashboard = self.client.get(reverse('home')).context['dashboard']['filas'][0]
        self.assertEqual(fila.resultado, detalle)
        self.assertEqual(fila.resultado, dashboard.resultado)
        self.assertEqual(fila.semaforo, dashboard.semaforo)
        metrica = next(p.metrica for p in fila.resultado.indicadores if p.indicador.pk == indicador.pk)
        self.assertEqual(metrica, calcular_resultado_indicador(indicador))

    def test_todas_las_unidades_reutilizan_9f_incluyendo_historial_temporal(self):
        for unidad in Indicador.Unidad.values:
            indicador = self.nuevo('20', unidad_medida=unidad,
                unidad_personalizada='árboles' if unidad == 'otra' else '')
            self.indicador_a = indicador
            self.ciclos([(self.instante(26, 9), self.instante(26, 14))])
        fila = self.fila_a()
        for ponderado in fila.resultado.indicadores:
            self.assertEqual(ponderado.metrica, calcular_resultado_indicador(ponderado.indicador))

    def test_observaciones_de_duracion_insuficiente_son_visibles(self):
        indicador = self.nuevo('100', unidad_medida='horas', unidad_personalizada='')
        self.vincular('cerrada', AporteIndicador.Estado.CUMPLE, indicador=indicador)
        self.assertContains(self.consultar(), 'sin historial temporal')
        self.assertIsNone(self.fila_a().resultado.cumplimiento_normalizado)

    def test_naranjo_reutiliza_umbral_del_dashboard(self):
        self.evaluado('100', '60')
        self.assertEqual(self.fila_a().semaforo.estado, 'naranjo')

    def test_limites_de_calendario_y_agregacion_global_explicitos(self):
        respuesta = self.consultar()
        for texto in ('no descuenta feriados', 'calendario inclusivo', 'La regla especial de pendientes en tubo',
                      'No se agrega un porcentaje global entre metas'):
            self.assertContains(respuesta, texto)

    def test_consulta_html_y_excel_no_escribe_en_base_ni_auditoria(self):
        self.evaluado('100', '120')
        for parametros in ({}, {'export': 'xlsx'}):
            with self.subTest(parametros=parametros), CaptureQueriesContext(connection) as consultas:
                self.consultar(**parametros)
            escrituras = [q['sql'] for q in consultas if q['sql'].lstrip().split()[0].upper()
                          in ('INSERT', 'UPDATE', 'DELETE', 'REPLACE')]
            self.assertEqual(escrituras, [])

    def test_consultas_no_crecen_por_metas_o_indicadores_temporales(self):
        indicador = self.nuevo('100', unidad_medida='horas', unidad_personalizada='')
        self.vincular('cerrada', AporteIndicador.Estado.CUMPLE, indicador=indicador)
        with CaptureQueriesContext(connection) as antes:
            self.consultar()
        for numero in range(4):
            meta = self.crear_meta(nombre=f'Meta temporal informe {numero}')
            indicador = self.nuevo('100', meta=meta, unidad_medida='dias', unidad_personalizada='')
            self.vincular('cerrada', AporteIndicador.Estado.CUMPLE, indicador=indicador)
        with CaptureQueriesContext(connection) as despues:
            self.consultar()
        self.assertEqual(len(antes), len(despues))

    def test_detalle_reutilizado_conserva_edicion_y_navegacion_previa(self):
        indicador = self.nuevo('100')
        self.entrar('delegado')
        respuesta = self.client.get(self.meta_a.get_absolute_url())
        self.assertContains(respuesta, reverse('sgr:indicador_editar', args=[indicador.pk]))
        for ruta in ('sgr:informe', 'sgr:periodo_lista', 'sgr:meta_lista', 'sgr:indicador_lista'):
            self.assertContains(respuesta, f'href="{reverse(ruta)}"')

    def test_excel_tiene_metadatos_y_valores_coherentes_con_html(self):
        self.evaluado('100', '231')
        libro = self.libro()
        self.assertEqual(libro.sheetnames, ['Metas', 'Indicadores'])
        hoja = libro['Metas']
        self.assertEqual(hoja['B1'].value, self.periodo.nombre)
        self.assertEqual(hoja['B3'].value, 'Global: todas las delegaciones')
        fila = next(f for f in hoja.iter_rows(min_row=8, values_only=True) if f[0] == self.meta_a.nombre)
        self.assertEqual(fila[2], 231)
        self.assertEqual(fila[3], 231)
        self.assertEqual(fila[4], 100)
        self.assertIn('Verde', fila[12])
        self.assertEqual(hoja.freeze_panes, 'C8')

    def test_excel_territorial_no_contiene_metas_o_indicadores_de_otra_delegacion(self):
        self.entrar('verificador')
        libro = self.libro(delegacion=self.b.pk, meta=self.meta_b.pk)
        valores = [c.value for hoja in libro for fila in hoja for c in fila]
        self.assertIn(self.meta_a.nombre, valores)
        self.assertNotIn(self.meta_b.nombre, valores)
        self.assertNotIn(self.indicador_b.nombre, valores)

    def test_excel_preserva_cero_none_y_sin_limite(self):
        cero = self.nuevo('30', unidad_medida='cantidad', unidad_personalizada='')
        limite = self.nuevo('30', sentido='menor', unidad_medida='cantidad', unidad_personalizada='')
        sin_datos = self.evaluado('40', None)
        filas = {f[2]: f for f in self.libro()['Indicadores'].iter_rows(min_row=8, values_only=True)}
        self.assertEqual(filas[cero.nombre][7:10], (0, 0, 'calculado'))
        self.assertEqual(filas[limite.nombre][7:10], (0, None, 'sin_limite'))
        self.assertEqual(filas[sin_datos.nombre][7:10], (None, None, 'sin_resultado'))

    def test_excel_no_interpreta_nombres_como_formulas(self):
        meta = self.crear_meta(nombre='=1+1')
        self.nuevo('100', meta=meta, nombre='=2+2')
        libro = self.libro()
        celdas = [c for hoja in libro for fila in hoja for c in fila]
        self.assertTrue(any(c.value == '=1+1' and c.data_type == 's' for c in celdas))
        self.assertTrue(any(c.value == '=2+2' and c.data_type == 's' for c in celdas))
        self.assertFalse(any(c.data_type == 'f' for c in celdas))

    def test_excel_reutiliza_periodo_seleccionado_y_excluye_los_demas(self):
        meta = self.crear_meta(periodo=self.otro_periodo, nombre='Meta exportada 2027')
        libro = self.libro(periodo=self.otro_periodo.pk)
        self.assertEqual(libro['Metas']['B1'].value, self.otro_periodo.nombre)
        valores = [fila[0] for fila in libro['Metas'].iter_rows(min_row=8, values_only=True)]
        self.assertEqual(valores, [meta.nombre])

    def test_exportacion_de_informe_sin_metas_conserva_encabezados(self):
        Meta.objects.filter(periodo=self.periodo).update(activo=False)
        libro = self.libro()
        self.assertEqual(libro['Metas'].max_row, 7)
        self.assertEqual(libro['Indicadores'].max_row, 7)

    def test_serializar_excel_no_hace_consultas_adicionales(self):
        self.evaluado('100', '120')
        informe = self.informe()
        with self.assertNumQueries(0):
            respuesta = exportar_informe_excel(informe)
        self.assertGreater(len(respuesta.content), 0)

    def test_rechaza_post_para_mantener_reporte_de_solo_lectura(self):
        self.assertEqual(self.client.post(reverse('sgr:informe'), {'periodo': self.periodo.pk}).status_code, 405)
