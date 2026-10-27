"""XLSX de valores ya calculados, sin fórmulas ni nuevas consultas."""
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font

LIMITACIONES = (
    'Resultados vigentes al consultar; no es una reconstrucción histórica. '
    'Calendario inclusivo; sin descontar feriados, licencias, vacaciones, emergencias ni compensatorios. '
    'Los indicadores inactivos no se agregan. Los valores no calculables quedan vacíos con estado explícito. '
    'Los porcentajes se expresan en escala 0 a 100, con sobrecumplimiento permitido. '
    'El semáforo representa solo la cobertura calculable; la regla especial del tubo queda fuera.'
)


def _agregar_fila(hoja, valores):
    hoja.append(valores)
    # Los nombres y observaciones son texto, incluso si comienzan con '='.
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = 's'


def exportar_informe_excel(informe):
    libro = Workbook()
    metas = libro.active
    metas.title = 'Metas'
    indicadores = libro.create_sheet('Indicadores')
    periodo = informe['periodo']
    for hoja in (metas, indicadores):
        for fila in (
            ['Informe SGR', periodo.nombre],
            ['Período', periodo.fecha_inicio, periodo.fecha_fin, 'Activo' if periodo.activo else 'Inactivo'],
            ['Contexto territorial', str(informe['contexto_territorial'])],
            ['Fecha de consulta', informe['fecha_corte']],
            ['Limitaciones', LIMITACIONES],
        ):
            _agregar_fila(hoja, fila)
        hoja.append([])
        hoja.freeze_panes = 'C8'
    _agregar_fila(metas, ['Meta', 'Delegación', 'Cumplimiento normalizado (%)',
        'Resultado ponderado bruto (pp)', 'Cobertura calculable (%)', 'Ponderación asignada (%)',
        'Ponderación pendiente (%)', 'Configuración completa', 'Indicadores sin ponderación',
        'Indicadores sin cumplimiento numérico', 'Indicadores sin límite', 'Avance esperado (%)', 'Semáforo'])
    _agregar_fila(indicadores, ['Meta', 'Delegación', 'Indicador', 'Activo', 'Unidad', 'Sentido',
        'Objetivo', 'Resultado actual', 'Cumplimiento (%)', 'Estado de cumplimiento',
        'Ponderación (%)', 'Resultado ponderado (pp)', 'Estado de ponderación', 'Base utilizada',
        'Sin valor validado', 'Sin duración calculable', 'Observaciones'])
    for fila in informe['filas']:
        resultado = fila.resultado
        _agregar_fila(metas, [fila.meta.nombre, str(fila.meta.delegacion),
            resultado.cumplimiento_normalizado, resultado.resultado_ponderado_bruto,
            resultado.cobertura_calculable, resultado.ponderacion_asignada, resultado.ponderacion_pendiente,
            'Sí' if resultado.configuracion_completa else 'No', resultado.sin_ponderacion,
            resultado.sin_resultado, resultado.sin_limite, informe['esperado'], fila.semaforo.texto])
        for ponderado in resultado.indicadores:
            indicador, metrica = ponderado.indicador, ponderado.metrica
            _agregar_fila(indicadores, [fila.meta.nombre, str(fila.meta.delegacion), indicador.nombre,
                'Sí' if indicador.activo else 'No', indicador.unidad_display, metrica.sentido,
                metrica.objetivo, metrica.resultado, metrica.cumplimiento.porcentaje,
                metrica.cumplimiento.estado.value, ponderado.ponderacion, ponderado.resultado_ponderado,
                ponderado.estado, metrica.base_utilizada, metrica.sin_valor_validado,
                metrica.sin_duracion_calculable, ' '.join(metrica.observaciones)])
    for hoja in (metas, indicadores):
        hoja.auto_filter.ref = f'A7:{hoja.cell(hoja.max_row, hoja.max_column).coordinate}'
        for celda in hoja[7]:
            celda.font = Font(bold=True)
        for columna in ('A', 'B', 'C'):
            hoja.column_dimensions[columna].width = 32
    buffer = BytesIO()
    libro.save(buffer)
    libro.close()
    respuesta = HttpResponse(buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="informe_sgr_periodo_{periodo.pk}.xlsx"'
    return respuesta
