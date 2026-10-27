from contextlib import closing
from io import BytesIO
from pathlib import Path
import importlib.util
import os
from unittest.mock import patch

from django.conf import settings
from django.core.exceptions import ImproperlyConfigured
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from catalogos.models import TipoAtencion
from sgr.test_metas import DatosMetasMixin


class RevisionDependenciasTests(SimpleTestCase):
    def test_manifiesto_de_dependencias_es_utf8_sin_bytes_nulos(self):
        texto = (Path(settings.BASE_DIR) / 'requirements.txt').read_text(encoding='utf-8-sig')
        self.assertNotIn('\x00', texto)
        self.assertTrue(all('==' in linea for linea in texto.splitlines() if linea.strip()))


class RevisionClaveConfiguracionTests(SimpleTestCase):
    def cargar_configuracion(self, valor):
        ruta = Path(settings.BASE_DIR) / 'config' / 'settings.py'
        spec = importlib.util.spec_from_file_location('configuracion_revision_aislada', ruta)
        modulo = importlib.util.module_from_spec(spec)
        with patch('dotenv.load_dotenv'), patch.dict(os.environ, {'DJANGO_SECRET_KEY': valor}):
            spec.loader.exec_module(modulo)
        return modulo

    def test_configuracion_usa_clave_del_entorno(self):
        ficticia = 'valor-solo-de-prueba-no-real-' + 'x' * 60
        modulo = self.cargar_configuracion(ficticia)
        self.assertTrue(modulo.SECRET_KEY == ficticia, 'Debe utilizar la clave del entorno aislado de prueba.')

    def test_falta_de_clave_no_utiliza_un_secreto_fijo(self):
        with self.assertRaises(ImproperlyConfigured):
            self.cargar_configuracion('')

    def test_clave_debil_se_rechaza_sin_mostrar_valor(self):
        with self.assertRaises(ImproperlyConfigured) as error:
            self.cargar_configuracion('django-insecure-valor-ficticio-no-real')
        self.assertIn('DJANGO_SECRET_KEY', str(error.exception))
        self.assertNotIn('valor-ficticio', str(error.exception))


class RevisionExportacionesTests(DatosMetasMixin, TestCase):
    def test_catalogo_no_exporta_textos_de_usuario_como_formulas(self):
        TipoAtencion.objects.create(codigo='=1+1', nombre='=2+2', descripcion='=3+3')
        self.entrar('administrador')
        respuesta = self.client.get(reverse('catalogos:tipo_lista'), {'export': 'xlsx'})
        self.assertEqual(respuesta.status_code, 200)
        with closing(load_workbook(BytesIO(respuesta.content))) as libro:
            celdas = [c for fila in libro.active for c in fila]
            for texto in ('=1+1', '=2+2', '=3+3'):
                celda = next(c for c in celdas if c.value == texto)
                self.assertEqual(celda.data_type, 's')

    def test_exportacion_conserva_filtro_y_estado_como_texto(self):
        TipoAtencion.objects.create(codigo='EXP-A', nombre='Nombre incluido')
        TipoAtencion.objects.create(codigo='EXP-B', nombre='Otro nombre')
        self.entrar('administrador')
        respuesta = self.client.get(reverse('catalogos:tipo_lista'), {'export': 'xlsx', 'q': 'incluido'})
        with closing(load_workbook(BytesIO(respuesta.content))) as libro:
            filas = list(libro.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][0:2], ('EXP-A', 'Nombre incluido'))
